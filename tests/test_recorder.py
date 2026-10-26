import json
import os
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer

from recorder.app import Recorder, make_handler, validate_config
from recorder.common import inside, cell_address
from recorder.compare import compare_files
from recorder.excel import CellCache, ExcelListener
from recorder.folders import FolderWorker
from recorder.store import Store

class Workspace(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)
        self.root = self.base/'watched'; self.root.mkdir()
        self.store = Store(self.base/'history')
        self.worker = FolderWorker(self.root, self.store, {'settle_seconds': 0, 'snapshot_limit_mb': 512})
    def tearDown(self):
        self.store.close(); self.temp.cleanup()
    def capture(self):
        self.worker.scan_once(); self.worker.scan_once()
    def kinds(self):
        return [e['event'] for e in self.store.recent(limit=1000)]

class FileCaptureTests(Workspace):
    def test_saved_versions_and_text_changes_preserve_original(self):
        p=self.root/'orders.csv';p.write_text('id,qty\nA,100\n')
        self.capture(); original=p.read_bytes()
        p.write_text('id,qty\nA,10\n');self.capture()
        versions=[e for e in self.store.recent() if e['event']=='file_version_captured']
        self.assertEqual(len(versions),2)
        self.assertEqual((self.store.directory/versions[-1]['snapshot']).read_bytes(), original)
        self.assertIn('snapshot_text_comparison',self.kinds())
        self.assertEqual(p.read_text(),'id,qty\nA,10\n')
    def test_other_binary_type_is_logged_not_copied(self):
        (self.root/'video.mp4').write_bytes(b'test')
        self.capture()
        self.assertIn('snapshot_not_selected',self.kinds())
        self.assertFalse(list(self.store.snapshots.rglob('*.mp4')))
    def test_document_snapshot_without_claiming_content_diff(self):
        p=self.root/'report.pdf';p.write_bytes(b'version1');self.capture()
        p.write_bytes(b'version2');self.capture()
        self.assertIn('comparison_not_available',self.kinds())
        self.assertEqual(len(list(self.store.snapshots.rglob('*.pdf'))),2)
    def test_offline_root_never_causes_mass_deletions(self):
        (self.root/'a.txt').write_text('keep');self.capture()
        moved=self.base/'disconnected';self.root.rename(moved)
        self.worker.scan_once()
        self.assertIn('folder_unavailable',self.kinds())
        self.assertNotIn('file_absent_on_scan',self.kinds())
        moved.rename(self.root);self.worker.scan_once()
        self.assertIn('folder_reconnected',self.kinds())
    def test_partial_enumeration_is_not_deletion(self):
        (self.root/'a.txt').write_text('keep');self.capture()
        with patch('recorder.folders.walk_files',side_effect=OSError('permission denied')):
            self.worker.scan_once()
        self.assertNotIn('file_absent_on_scan',self.kinds())
    def test_deleted_file_keeps_snapshot(self):
        p=self.root/'a.txt';p.write_text('keep');self.capture();p.unlink();self.worker.scan_once()
        self.assertIn('file_absent_on_scan',self.kinds())
        self.assertEqual(len(list(self.store.snapshots.rglob('*.txt'))),1)
    def test_source_changed_during_copy_does_not_commit_snapshot(self):
        p=self.root/'a.txt';p.write_text('keep')
        from recorder.folders import signature
        sig=signature(p)
        with patch('recorder.folders.signature',return_value=[9,9,9,9]):
            with self.assertRaises(OSError):self.worker.capture(str(p),sig)
        self.assertFalse(list(self.store.snapshots.rglob('*.partial')))
        self.assertFalse(list(self.store.snapshots.rglob('*.txt')))
    def test_locked_file_is_retried_without_replacing_baseline(self):
        p=self.root/'a.txt';p.write_text('old');self.capture()
        before=self.store.inventory(self.worker.root)
        p.write_text('new');self.worker.scan_once()
        with patch.object(self.worker,'capture',side_effect=PermissionError('locked')):
            self.worker.scan_once()
        self.assertEqual(self.store.inventory(self.worker.root),before)
        self.assertIn('snapshot_retry',self.kinds())
        self.worker.scan_once()
        self.assertNotEqual(self.store.inventory(self.worker.root),before)
    def test_temporary_files_are_ignored(self):
        for name in ['~$orders.xlsx','incomplete.tmp']:(self.root/name).write_text('ignored')
        self.capture();self.assertFalse(self.store.inventory(self.worker.root))
    def test_no_duplicate_capture_after_unchanged_restart(self):
        (self.root/'a.txt').write_text('a');self.capture()
        fresh=FolderWorker(self.root,self.store,{'settle_seconds':0})
        fresh.scan_once();fresh.scan_once()
        self.assertEqual(self.kinds().count('file_version_captured'),1)
    def test_paused_worker_does_not_capture(self):
        (self.root/'a.txt').write_text('a');self.worker.enabled=lambda:False
        self.capture();self.assertEqual(self.kinds(),[])
    def test_symlink_does_not_escape_scope(self):
        outside=self.base/'outside';outside.mkdir();(outside/'secret.txt').write_text('secret')
        try:(self.root/'link').symlink_to(outside,target_is_directory=True)
        except OSError:self.skipTest('Symlink privilege unavailable')
        self.capture();self.assertFalse(self.store.inventory(self.worker.root))

class ComparisonTests(Workspace):
    def test_xlsx_values_formulas_and_sheet_changes(self):
        import openpyxl
        before=self.base/'before.xlsx';after=self.base/'after.xlsx'
        wb=openpyxl.Workbook();wb.active['A1']=100;wb.active['B1']='=A1*2';wb.save(before)
        wb.active['A1']=10;wb.active['B1']='=A1*3';wb.create_sheet('New');wb.save(after);wb.close()
        original=before.read_bytes()
        compare_files(before,after,self.store.emit,'source.xlsx')
        e=next(e for e in self.store.recent() if e['event']=='snapshot_cell_comparison')
        self.assertEqual(e['total_changes'],2);self.assertFalse(e['details_truncated'])
        self.assertIn('snapshot_sheet_added',self.kinds());self.assertEqual(before.read_bytes(),original)
    def test_comparison_limit_is_explicit(self):
        import openpyxl
        paths=[self.base/'a.xlsx',self.base/'b.xlsx']
        for p in paths:
            wb=openpyxl.Workbook();wb.active['B2']=1;wb.save(p);wb.close()
        with self.assertRaises(ValueError):compare_files(*paths,self.store.emit,'source',max_cells=1)

class CacheTests(unittest.TestCase):
    def test_unknown_old_is_not_invented(self):
        c=CellCache();e=c.change('sheet',{'A1':{'value':10,'formula':None}})[0]
        self.assertFalse(e['old_value_known']);self.assertIsNone(e['old_value'])
    def test_old_and_new_formulas_are_distinct(self):
        c=CellCache();c.seed('sheet',{'A1':{'value':2,'formula':'=1+1'}})
        e=c.change('sheet',{'A1':{'value':3,'formula':'=1+2'}})[0]
        self.assertTrue(e['old_value_known']);self.assertEqual(e['old_formula'],'=1+1');self.assertEqual(e['new_formula'],'=1+2')
        c.invalidate('sheet');self.assertFalse(c.change('sheet',{'A1':{'value':4}})[0]['old_value_known'])
    def test_cell_addresses(self):
        self.assertEqual(cell_address(2,28),'AB2')

class StoreTests(Workspace):
    def test_concurrent_events_are_complete_json(self):
        def write(n):
            for i in range(10):self.store.emit('concurrent',thread=n,item=i)
        ts=[threading.Thread(target=write,args=(n,)) for n in range(4)]
        for t in ts:t.start()
        for t in ts:t.join()
        self.assertEqual(len(self.store.recent()),40)
        lines=list(self.store.logs.glob('*.jsonl'))[0].read_text().splitlines()
        self.assertEqual(len({json.loads(x)['id'] for x in lines}),40)
    def test_literal_search_does_not_expand_sql_wildcards(self):
        self.store.emit('file','100%_done');self.store.emit('other','100Xdone')
        self.assertEqual(len(self.store.recent(search='%_')),1)
    def test_jsonl_rebuilds_lost_index_entry(self):
        event=self.store.emit('recoverable','x')
        with self.store.lock:self.store.db.execute('DELETE FROM events');self.store.db.commit()
        recovered=Store(self.store.directory)
        try:self.assertEqual(recovered.recent()[0]['id'],event['id'])
        finally:recovered.close()

class ConfigurationTests(Workspace):
    @unittest.skipUnless(os.name == 'nt', 'Windows short-path alias check')
    def test_windows_short_path_cannot_bypass_history_exclusion(self):
        import win32api
        short=win32api.GetShortPathName(str(self.base))
        with self.assertRaises(ValueError):validate_config({'roots':[short]},self.store.directory)
    def test_reject_output_within_monitored_root(self):
        with self.assertRaises(ValueError):validate_config({'roots':[str(self.base)]},self.store.directory)
    def test_reject_nested_roots(self):
        with self.assertRaises(ValueError):validate_config({'roots':[str(self.root),str(self.root/'sub')]},self.store.directory)
    def test_accept_unavailable_absolute_network_mount(self):
        config=validate_config({'roots':[str(self.base/'offline')]},self.store.directory)
        self.assertEqual(len(config['roots']),1)
    def test_reject_invalid_intervals(self):
        with self.assertRaises(ValueError):validate_config({'scan_seconds':0},self.store.directory)
    def test_scope_blocks_sibling_prefix(self):
        self.assertFalse(inside(self.base/'watched-other',self.root))

class HttpTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.recorder=Recorder(Path(self.temp.name)/'history')
        self.server=ThreadingHTTPServer(('127.0.0.1',0),make_handler(self.recorder,'secret',lambda:None))
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
        self.url=f'http://127.0.0.1:{self.server.server_port}'
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.recorder.store.close();self.temp.cleanup()
    def request(self,path,body=None,token='secret',host=None):
        headers={'X-Recorder-Token':token,'Content-Type':'application/json'}
        if host:headers['Host']=host
        return urllib.request.urlopen(urllib.request.Request(self.url+path,headers=headers,data=json.dumps(body).encode() if body is not None else None),timeout=3)
    def test_no_token_cannot_read_content(self):
        with self.assertRaises(urllib.error.HTTPError) as e:self.request('/api/events',token='')
        self.assertEqual(e.exception.code,403)
    def test_untrusted_host_is_rejected(self):
        with self.assertRaises(urllib.error.HTTPError) as e:self.request('/api/status',host='evil.example')
        self.assertEqual(e.exception.code,403)
    def test_note_roundtrip_and_export(self):
        text='<script>alert(1)</script> — only Q3'
        with self.request('/api/note',{'text':text,'event_id':'abc','scope':'case_note'}) as r:self.assertEqual(r.status,200)
        with self.request('/api/events') as r:data=json.load(r)
        self.assertEqual(data[0]['text'],text);self.assertTrue(data[0]['reusable_rule_not_approved'])
        with self.request('/api/export') as r:self.assertIn('only Q3',r.read().decode())
    def test_note_empty_rejected(self):
        with self.assertRaises(urllib.error.HTTPError) as e:self.request('/api/note',{'text':''})
        self.assertEqual(e.exception.code,400)
    def test_page_has_no_external_scripts(self):
        with self.request('/') as r:
            html=r.read().decode();self.assertIn("textContent=content",html);self.assertNotIn('<script src=',html)
            self.assertEqual(r.headers['Referrer-Policy'],'no-referrer')

class ExcelEventTests(Workspace):
    def make_sheet(self, value=100, formula=100, path=None):
        from types import SimpleNamespace as N
        area=N(Row=1,Column=1,Rows=N(Count=1),Columns=N(Count=1),Value2=value,Formula=formula)
        target=N(CountLarge=1,Areas=[area],Address='$A$1')
        wb=N(Path=str(self.root),FullName=str(path or self.root/'book.xlsb'))
        sh=N(Parent=wb,CodeName='Sheet1',Name='Orders',UsedRange=target)
        wb.Worksheets=[sh]
        return wb,sh,target,area
    def listener(self):
        return ExcelListener(self.store,{'excel_cache_cells':100,'excel_event_cells':10},lambda:[str(self.root)])
    def test_live_event_records_observed_old_new_and_unsaved_status(self):
        listener=self.listener();wb,sh,target,area=self.make_sheet()
        listener.opened(wb,'instance');area.Value2=10;area.Formula=10
        listener.handlers('instance')().OnSheetChange(sh,target)
        event=self.store.recent()[0]
        self.assertEqual(event['event'],'excel_range_changed')
        self.assertEqual(event['changes'][0]['old_value'],100)
        self.assertEqual(event['changes'][0]['new_value'],10)
        self.assertEqual(event['save_state'],'not_confirmed_saved')
    def test_outside_workbook_content_is_not_logged(self):
        listener=self.listener();wb,sh,target,area=self.make_sheet(path=self.base/'private.xlsx')
        listener.changed(sh,target,'instance');self.assertEqual(self.kinds(),[])
    def test_oversized_live_event_is_explicit_and_invalidates_old_cache(self):
        listener=self.listener();wb,sh,target,area=self.make_sheet();listener.prime(wb,'instance')
        target.CountLarge=100000;listener.changed(sh,target,'instance')
        e=self.store.recent()[0];self.assertFalse(e['details_captured']);self.assertTrue(e['old_baseline_invalidated'])
    def test_reopened_workbook_does_not_reuse_discarded_edit_baseline(self):
        listener=self.listener();wb,sh,target,area=self.make_sheet();listener.opened(wb,'instance')
        area.Value2=10;area.Formula=10;listener.changed(sh,target,'instance')
        area.Value2=100;area.Formula=100;listener.opened(wb,'instance')
        area.Value2=20;area.Formula=20;listener.changed(sh,target,'instance')
        self.assertEqual(self.store.recent()[0]['changes'][0]['old_value'],100)
    def test_save_request_and_result_have_separate_semantics(self):
        listener=self.listener();wb,sh,target,area=self.make_sheet();events=listener.handlers('instance')()
        events.OnWorkbookBeforeSave(wb,True,False);events.OnWorkbookAfterSave(wb,False)
        logs=self.store.recent();self.assertFalse(next(e for e in logs if e['event']=='excel_save_result')['success'])
        self.assertTrue(next(e for e in logs if e['event']=='excel_save_requested')['success_not_implied'])

if __name__=='__main__':unittest.main()
